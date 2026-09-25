"""
29_evidence_grading.py — Independent grading of the evidence for each registry
intervention by two models from different developers, alongside the
investigators' grade, with a physician adjudication sheet for disagreements.

Each grader sees, for each intervention, its name, the cause and the stage of the
pathway to death it is meant to act on, and a neutral summary of the strongest
available evidence (citation, design, outcome, effect estimate). Sentences in
which the investigators stated their own modeling decisions are removed, and no
grader sees how often committees requested any intervention. Graders may also
draw on their own knowledge of the literature.

Scale, for the effect on the named stage of the pathway to pregnancy-related death:
  Demonstrated benefit  randomized evidence or a synthesis of randomized trials
                        showing benefit on this outcome or a direct precursor of it
  Moderate              consistent nonrandomized or indirect evidence of benefit
  Weak                  limited, confounded, or extrapolated evidence
  Null                  tested for this or a directly relevant outcome, no benefit
  Absent                not tested for maternal outcomes
  Not applicable        not an intervention on a pathway (for example surveillance)

Output: results/evidence_grades.json, validation/MMRC_evidence_grading_adjudication.xlsx,
        results/evidence_grading_key.csv
"""
from __future__ import annotations
import csv, importlib.util, json, random, re, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RES, OUT = ROOT / "results", ROOT / "validation"
GRADES = ["Demonstrated benefit", "Moderate", "Weak", "Null", "Absent", "Not applicable"]
MODELS = ["claude-opus-5-5", "gpt-6-astra"]
STAGE = {"incidence": "whether the complication occurs",
         "smm": "whether the complication progresses to severe maternal morbidity",
         "cfr": "whether severe morbidity ends in death",
         "exposure": "the length of the postpartum period in which care can reach a person",
         "none": "no specific stage (a general program)"}
DECISION = re.compile(r"(Entered at[^.]*\.|[^.]*labeled as such[^.]*\.|[^.]*weakest link[^.]*\.|"
                      r"A null result from[^.]*\.|[^.]*is set well below[^.]*\.|"
                      r"In the model[^.]*\.|[^.]*with an upward band[^.]*\.)", re.I)


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def canon_grade(g: str) -> str:
    g = g.lower()
    if "absent for mortality" in g:
        return "Absent"
    for pre, lab in (("demonstrated", "Demonstrated benefit"), ("moderate", "Moderate"),
                     ("weak", "Weak"), ("null", "Null"), ("absent", "Absent")):
        if g.startswith(pre):
            return lab
    return "Not applicable"


PROMPT = """You are grading the strength of evidence for interventions intended to prevent
pregnancy-related death in the United States. For each intervention, grade the evidence
that it benefits the named stage of the pathway to death, using this scale:

Demonstrated benefit: randomized evidence or a synthesis of randomized trials showing
  benefit on this outcome or a direct clinical precursor of it.
Moderate: consistent nonrandomized or indirect evidence of benefit.
Weak: limited, confounded, or extrapolated evidence.
Null: tested for this or a directly relevant outcome, with no benefit.
Absent: not tested for maternal outcomes.
Not applicable: not an intervention acting on a clinical pathway (for example, data systems).

Use the evidence summary provided and your own knowledge of the literature. Grade each
intervention independently. Return only a JSON array:
[{{"key": "<key>", "grade": "<one grade>", "reason": "<one sentence>"}}, ...]

Interventions:
{items}"""


def main() -> None:
    iv = _load("iv", "12_interventions.py"); dm = _load("dm", "25_dual_model_classify.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    ints = iv.build(P)
    items = []
    for i in ints:
        ev = DECISION.sub("", i.source).strip()
        items.append(f'- key: {i.key}\n  intervention: {i.label}\n  cause of death: '
                     f'{", ".join(i.causes) or "not cause-specific"}\n  stage: {STAGE.get(i.stage, i.stage)}\n'
                     f'  evidence: {ev}')
    prompt = PROMPT.format(items="\n".join(items))

    out = {"investigator": {i.key: canon_grade(i.grade) for i in ints}, "reasons": {}}
    for m in MODELS:
        call = dm.make_caller(m)
        res = dm.with_retry(call, prompt, dm.parse_array)
        got = {r["key"]: r for r in res if r.get("grade") in GRADES}
        out[m] = {k: got[k]["grade"] for k in got}
        out["reasons"][m] = {k: got[k].get("reason", "") for k in got}
        print(m, len(got), "graded")
    out["labels"] = {i.key: i.label for i in ints}
    out["prompt"] = prompt
    json.dump(out, open(RES / "evidence_grades.json", "w"), indent=1)

    # agreement
    keys = [i.key for i in ints]
    trip = {k: (out["investigator"][k], out[MODELS[0]].get(k), out[MODELS[1]].get(k)) for k in keys}
    unanimous = [k for k, t in trip.items() if len(set(t)) == 1]
    print(f"unanimous {len(unanimous)} of {len(keys)}")

    # adjudication sheet for non-unanimous interventions, options unlabeled and shuffled
    rng = random.Random(20260927)
    wb = Workbook(); ws = wb.active; ws.title = "Evidence grades"
    cols = [("Row", 6), ("Intervention", 34), ("Cause of death", 24), ("Stage of the pathway", 30),
            ("Evidence summary", 70), ("Grade A", 18), ("Grade B", 18), ("Grade C", 18),
            ("Your grade", 18), ("Comment", 28)]
    fill = PatternFill("solid", fgColor="DDE6EE")
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    dv = DataValidation(type="list", formula1='"' + ",".join(GRADES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    key_rows = []
    r = 2
    for i in ints:
        if i.key in unanimous:
            continue
        srcs = [("investigator", trip[i.key][0]), (MODELS[0], trip[i.key][1]), (MODELS[1], trip[i.key][2])]
        rng.shuffle(srcs)
        vals = [r - 1, i.label, ", ".join(i.causes) or "not cause-specific", STAGE.get(i.stage, i.stage),
                DECISION.sub("", i.source).strip()] + [g for _, g in srcs]
        for j, v in enumerate(vals, 1):
            ws.cell(row=r, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        dv.add(f"I{r}")
        key_rows.append([r, i.key] + [s for s, _ in srcs])
        r += 1
    ws2 = wb.create_sheet("Adjudicator"); ws2["A1"] = "Initials"; ws2["A2"] = "Date completed"
    wb.save(OUT / "MMRC_evidence_grading_adjudication.xlsx")
    with open(RES / "evidence_grading_key.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["excel_row", "key", "grade_A_source", "grade_B_source", "grade_C_source"])
        w.writerows(key_rows)
    print(f"adjudication rows: {r - 2}")


if __name__ == "__main__":
    main()
