"""
27_adjudication_sheet.py — The disagreements between the two model families, for a
physician to resolve, ordered by whether they can change a result.

A recommendation enters the workbook if the models disagree on cause, policy
lever, requested intervention, or whether the actor named is a binding actor.
Items are split into two tiers. Priority items can change a headline result: an
intervention disagreement in which either label is an intervention with evidence
of benefit (the evidence result and the model), and a cause disagreement that
would flip a state's silence status on a documented cause (the silence result).
Committee-written cause labels on which the models disagree are also priority,
because they define documented burden. All other disagreements change only
descriptive percentages and are optional.

The two labels appear as Option 1 and Option 2 in random order, so the adjudicator
does not know which model gave which. This step is adjudication with labels
visible and is reported as such; the blinded physician sample
(24_validation_packet.py) is the reference standard.

Output: validation/MMRC_adjudication.xlsx, results/adjudication_key.csv,
        results/adjudication_tiers.json
"""
from __future__ import annotations
import collections, csv, importlib.util, json, random, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RES, OUT = ROOT / "results", ROOT / "validation"
SRC = Path.home() / "waymark-local/notebooks/dark-health-data/data/processed/mmrc"
SEED = 20260926
BINDING = {"Legislature or statute", "Medicaid or state health agency",
           "Managed care organization or payer"}
FILL = PatternFill("solid", fgColor="DDE6EE")


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def jl(name, field=None):
    rows = {json.loads(l)["id"]: json.loads(l) for l in open(RES / name)}
    return {k: (v[field] if field else v) for k, v in rows.items()}


def header(ws, cols):
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w


def main() -> None:
    lc = _load("lc", "07_llm_classify.py"); iv = _load("iv", "12_interventions.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    ints = iv.build(P)
    lab = {i.key: i.label for i in ints}; lab["none"] = "No specific intervention"
    eff = {i.key for i in ints if i.grade.lower().startswith(("demonstrated", "moderate"))
           and "absent for mortality" not in i.grade.lower()}
    recs = {r["id"]: r for r in lc.load_recs()}
    A, B = jl("llm_labels_passA.jsonl"), jl("llm_labels_passB.jsonl")
    IA, IB = jl("intervention_map_passA.jsonl", "intervention"), jl("intervention_map_passB.jsonl", "intervention")

    # silence flips: burden from the primary cause-label map
    cmap = json.load(open(RES / "cause_label_map.json"))["map"]
    burden = collections.defaultdict(set)
    for f in csv.DictReader(open(SRC / "mmrc_findings.csv")):
        d = cmap.get(f["leading_cause"].strip())
        if d and d != "Not a cause of death":
            burden[f["state"]].add(d)
    cnt = collections.Counter((recs[i]["state"], A[i]["cause_domain"]) for i in A)

    def flips(i):
        a, b, s = A[i]["cause_domain"], B[i]["cause_domain"], recs[i]["state"]
        return (a in burden[s] and cnt[(s, a)] == 1) or (b in burden[s] and cnt[(s, b)] == 0)

    priority, remaining = [], []
    for i in sorted(set(A) & set(B) & set(IA) & set(IB)):
        items = []
        if A[i]["cause_domain"] != B[i]["cause_domain"]:
            items.append(("Cause of death addressed", A[i]["cause_domain"], B[i]["cause_domain"],
                          flips(i)))
        if A[i]["policy_lever"] != B[i]["policy_lever"]:
            items.append(("Policy lever", A[i]["policy_lever"], B[i]["policy_lever"], False))
        if IA[i] != IB[i]:
            items.append(("Specific intervention requested", lab[IA[i]], lab[IB[i]],
                          IA[i] in eff or IB[i] in eff))
        ba = "Yes" if A[i]["actor_authority"] in BINDING else "No"
        bb = "Yes" if B[i]["actor_authority"] in BINDING else "No"
        if ba != bb:
            items.append(("Binding actor named", ba, bb, False))
        pri = [x for x in items if x[3]]
        rest = [x for x in items if not x[3]]
        if pri:
            priority.append((i, pri))
        if rest:
            remaining.append((i, rest))

    rng = random.Random(SEED)
    rng.shuffle(priority); rng.shuffle(remaining)
    key = []
    cols = [("Row", 6), ("State", 7), ("Report section or category", 26), ("Recommendation text", 64),
            ("Field", 24), ("Option 1", 30), ("Option 2", 30), ("Your decision", 14),
            ("If neither, correct label", 30), ("Comment", 24)]

    wb = Workbook()
    readme = wb.active; readme.title = "README"
    readme["A1"] = ("Please complete sheet 1 (Priority) and sheet 2 (Cause labels); these items can change "
                    "the paper's results. Sheet 3 (Remaining) is optional and affects only descriptive "
                    "percentages. For each row choose Option 1, Option 2, or Neither using the definitions "
                    "in CODING_INSTRUCTIONS.md. The two options are in random order and are not labeled by "
                    "source. If you are also a blinded coder, finish your blinded coding file first.")
    readme["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 110; readme.row_dimensions[1].height = 90

    def fill(ws, items, sheet_name):
        header(ws, cols)
        dv = DataValidation(type="list", formula1='"Option 1,Option 2,Neither"', allow_blank=True)
        ws.add_data_validation(dv)
        r = 2
        for n, (i, dis) in enumerate(items, 1):
            rec = recs[i]
            for name, a, b, _ in dis:
                a_first = rng.random() < 0.5
                o1, o2 = (a, b) if a_first else (b, a)
                for j, v in enumerate([n, rec["state"], rec["category"], rec["text"], name, o1, o2], 1):
                    ws.cell(row=r, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
                dv.add(f"H{r}")
                key.append([sheet_name, r, i, name, "primary" if a_first else "second",
                            "second" if a_first else "primary"])
                r += 1
        ws.freeze_panes = "E2"
        return r - 2

    n_pri = fill(wb.create_sheet("1 Priority"), priority, "1 Priority")

    ca = json.load(open(RES / "cause_label_map_claude-opus-5-5.json"))["map"]
    cb = json.load(open(RES / "cause_label_map_gpt-6-astra.json"))["map"]
    cdis = [s for s in sorted(set(ca) & set(cb)) if ca[s] != cb[s]]
    rng.shuffle(cdis)
    ws2 = wb.create_sheet("2 Cause labels")
    header(ws2, [("Row", 6), ("Committee-written cause label", 50), ("Option 1", 28), ("Option 2", 28),
                 ("Your decision", 14), ("If neither, correct category", 28), ("Comment", 24)])
    dv2 = DataValidation(type="list", formula1='"Option 1,Option 2,Neither"', allow_blank=True)
    ws2.add_data_validation(dv2)
    for n, s in enumerate(cdis, 2):
        a_first = rng.random() < 0.5
        o1, o2 = (ca[s], cb[s]) if a_first else (cb[s], ca[s])
        for j, v in enumerate([n - 1, s, o1, o2], 1):
            ws2.cell(row=n, column=j, value=v)
        dv2.add(f"E{n}")
        key.append(["2 Cause labels", n, s, "cause category", "primary" if a_first else "second",
                    "second" if a_first else "primary"])

    n_rem = fill(wb.create_sheet("3 Remaining (optional)"), remaining, "3 Remaining (optional)")
    ws4 = wb.create_sheet("Adjudicator")
    ws4["A1"] = "Initials"; ws4["A2"] = "Date completed"
    wb.save(OUT / "MMRC_adjudication.xlsx")

    with open(RES / "adjudication_key.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["sheet", "excel_row", "id_or_string", "field", "option1_source", "option2_source"])
        w.writerows(key)
    by = collections.Counter(x[0] for _, d in priority for x in d)
    json.dump({"priority_recommendations": len(priority), "priority_items": n_pri,
               "priority_by_field": dict(by), "cause_label_items": len(cdis),
               "remaining_recommendations": len(remaining), "remaining_items": n_rem},
              open(RES / "adjudication_tiers.json", "w"), indent=1)
    print(f"priority: {len(priority)} recommendations, {n_pri} items {dict(by)}; "
          f"cause labels: {len(cdis)}; optional remaining: {len(remaining)} recommendations, {n_rem} items")


if __name__ == "__main__":
    main()
