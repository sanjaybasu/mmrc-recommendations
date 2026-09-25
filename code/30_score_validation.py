"""
30_score_validation.py — Score the automated classification against the blinded
physician reference standard, and fix the final evidence grades.

Validation. The physician coded a stratified sample blind to both models' labels.
For each field we report Cohen kappa and simple agreement between the physician
and each model, and between the physician and the final adjudicated labels. The
sample was stratified to include every requested intervention, so rare categories
are overrepresented relative to the corpus; the statistics describe agreement on
the sample and are reported unweighted.

Evidence grades. Each intervention was graded by the investigators and by two
models; where the three disagreed, the physician's adjudicated grade is final.

Output: results/validation_scores.json, results/evidence_grades_final.json
"""
from __future__ import annotations
import csv, importlib.util, json, sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RES, VAL = ROOT / "results", ROOT / "validation"
FIELDS = [("Cause of death addressed", "cause_domain"), ("Policy lever", "policy_lever"),
          ("Most authoritative actor named", "actor_authority"), ("Coverage period", "coverage_period"),
          ("Specific intervention requested", "intervention")]
BINDING = {"Legislature or statute", "Medicaid or state health agency",
           "Managed care organization or payer"}


def kappa(a, b):
    cats = sorted(set(a) | set(b)); ix = {c: i for i, c in enumerate(cats)}
    m = np.zeros((len(cats), len(cats)))
    for x, y in zip(a, b):
        m[ix[x], ix[y]] += 1
    po = np.trace(m) / m.sum(); pe = float((m.sum(0) / m.sum()) @ (m.sum(1) / m.sum()))
    return (float((po - pe) / (1 - pe)) if pe < 1 else 1.0), float(po)


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def jl(path, field=None):
    rows = {json.loads(l)["id"]: json.loads(l) for l in open(path)}
    return {k: (v[field] if field else v) for k, v in rows.items()}


def main() -> None:
    iv = _load("iv", "12_interventions.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    lab = {i.key: i.label for i in iv._registry(P)}; lab["none"] = "No specific intervention"

    key = [r for r in csv.DictReader(open(RES / "validation_key.csv"))]
    rec_ids = {int(r["row"]): int(r["id_or_string"]) for r in key if r["sheet"] == "Recommendations"}
    cause_rows = {int(r["row"]): r["id_or_string"] for r in key if r["sheet"] == "Cause labels"}

    wb = load_workbook(VAL / "MMRC_blinded_coding_ADJ.xlsx", data_only=True)
    ws = wb["Recommendations"]; hdr = [c.value for c in ws[1]]
    phys = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        n = int(row[0]); i = rec_ids[n]
        phys[i] = {f: row[hdr.index(h)] for h, f in FIELDS}
        phys[i]["unsure"] = str(row[hdr.index("Unsure? (Y)")]).strip().upper() == "Y"

    sources = {
        "claude-opus-5-5": (jl(RES / "labels_claude-opus-5-5.jsonl"), jl(RES / "imap_claude-opus-5-5.jsonl", "intervention")),
        "gpt-6-astra": (jl(RES / "labels_gpt-6-astra.jsonl"), jl(RES / "imap_gpt-6-astra.jsonl", "intervention")),
        "final_adjudicated": (jl(RES / "llm_labels_passA.jsonl"), jl(RES / "intervention_map_passA.jsonl", "intervention")),
    }
    out = {"n_recommendations": len(phys), "n_unsure": sum(p["unsure"] for p in phys.values()),
           "fields": {}}
    for s, (L, I) in sources.items():
        out["fields"][s] = {}
        for _, f in FIELDS:
            ids = sorted(phys)
            ph = [phys[i][f] for i in ids]
            md = [lab[I[i]] if f == "intervention" else L[i][f] for i in ids]
            k, po = kappa(ph, md)
            out["fields"][s][f] = {"kappa": round(k, 3), "agreement": round(po, 3), "n": len(ids)}
        ph_b = [phys[i]["actor_authority"] in BINDING for i in sorted(phys)]
        md_b = [L[i]["actor_authority"] in BINDING for i in sorted(phys)]
        k, po = kappa(ph_b, md_b)
        out["fields"][s]["binding_actor"] = {"kappa": round(k, 3), "agreement": round(po, 3), "n": len(phys)}

    # committee-written cause labels
    ws2 = wb["Cause labels"]
    pc = {cause_rows[int(r[0])]: r[2] for r in ws2.iter_rows(min_row=2, values_only=True)}
    for s, f in (("claude-opus-5-5", RES / "cause_label_map_claude-opus-5-5.json"),
                 ("gpt-6-astra", RES / "cause_label_map_gpt-6-astra.json"),
                 ("final_adjudicated", RES / "cause_label_map.json")):
        m = json.load(open(f))["map"]
        k, po = kappa([pc[x] for x in pc], [m[x] for x in pc])
        out["fields"][s]["cause_label"] = {"kappa": round(k, 3), "agreement": round(po, 3), "n": len(pc)}
    json.dump(out, open(RES / "validation_scores.json", "w"), indent=1)
    for s, v in out["fields"].items():
        print(s, {f: (x["kappa"], x["agreement"]) for f, x in v.items()})

    # final evidence grades
    eg = json.load(open(RES / "evidence_grades.json"))
    gk = {int(r["excel_row"]): r for r in csv.DictReader(open(RES / "evidence_grading_key.csv"))}
    wb2 = load_workbook(VAL / "MMRC_evidence_grading_adjudication_ADJ.xlsx", data_only=True)
    ws3 = wb2["Evidence grades"]; hdr3 = [c.value for c in ws3[1]]
    adj = {}
    for row in ws3.iter_rows(min_row=2):
        g = row[hdr3.index("Your grade")].value
        if g:
            adj[gk[row[0].row]["key"]] = g
    final, src = {}, {}
    for k in eg["investigator"]:
        trip = {eg["investigator"][k], eg["claude-opus-5-5"].get(k), eg["gpt-6-astra"].get(k)}
        if k in adj:
            final[k], src[k] = adj[k], "physician adjudication"
        elif len(trip) == 1:
            final[k], src[k] = eg["investigator"][k], "unanimous"
        else:
            raise SystemExit(f"{k}: graders disagree and no adjudicated grade")
    agree = lambda a, b: sum(eg[a][k] == eg[b].get(k) for k in final) / len(final)
    summary = {"final": final, "source": src,
               "n_interventions": len(final),
               "n_unanimous": sum(v == "unanimous" for v in src.values()),
               "n_adjudicated": len(adj),
               "adjudicated_changed_investigator_grade": sum(final[k] != eg["investigator"][k] for k in adj),
               "pairwise_agreement": {"investigator_vs_opus": agree("investigator", "claude-opus-5-5"),
                                      "investigator_vs_gpt": agree("investigator", "gpt-6-astra"),
                                      "opus_vs_gpt": agree("claude-opus-5-5", "gpt-6-astra")}}
    json.dump(summary, open(RES / "evidence_grades_final.json", "w"), indent=1)
    print({k: v for k, v in summary.items() if k not in ("final", "source")})


if __name__ == "__main__":
    main()
