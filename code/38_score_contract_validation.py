"""
38_score_contract_validation.py — Accuracy of the text rules that decide whether a
Medicaid contract names a cause of pregnancy-related death in a maternal context,
and whether it ties that cause to a payment consequence.

Reference standard. An investigator (K.O.W.) read 80 passages per rule (script 35;
40 flagged and 40 not flagged by the rule, shuffled, rule decision hidden) and
judged each Yes, No, or Unclear.

Statistics. Positive predictive value among flagged passages and negative
predictive value among unflagged passages, each with a Wilson 95% CI; Unclear
judgments are excluded and, as a sensitivity analysis, counted against the rule.
Because the flagged and unflagged passages were sampled separately, sensitivity
and specificity are not estimated.

Narrower windows. Each passage is centered on the cause term (350 characters each
side for the cause rule, 600 for the payment rule), so the distance from the cause
term to the nearest maternal or payment term can be measured within the passage.
Positive predictive value is recomputed for flagged passages that also meet a
narrower window (150 characters for the cause rule, 500 for the payment rule).
Distances are measured on whitespace-collapsed text and are approximate.

Output: results/contract_validation_scores.json
"""
from __future__ import annotations
import csv, importlib.util, json, math, sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RES, VAL = ROOT / "results", ROOT / "validation"
WB = VAL / "MMRC_contract_text_validation_KW.xlsx"
NARROW = {"Cause mentions": 150, "Payment provisions": 500}


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def wilson(k, n, z=1.96):
    if n == 0:
        return None
    p = k / n; d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d; h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [round(max(0, c - h), 3), round(min(1, c + h), 3)]


def nearest(text, cause_pat, other_pat):
    """Distance from the cause match nearest the passage center to the nearest other term."""
    mid = len(text) / 2
    cm = min(cause_pat.finditer(text), key=lambda m: abs(m.start() - mid), default=None)
    if cm is None:
        return None
    pos = [m.start() for m in other_pat.finditer(text)]
    return min((abs(p - cm.start()) for p in pos), default=None)


def main() -> None:
    chain = _load("chain", "04_accountability_chain.py")
    import re
    cause_pats = {d: re.compile(p, re.I) for d, p in chain.CAUSE_TERMS.items()}
    other = {"Cause mentions": chain.MATERNAL_CONTEXT, "Payment provisions": chain.WITHHOLD}
    key = {(r["sheet"], int(r["row"])): r for r in csv.DictReader(open(RES / "contract_validation_key.csv"))}
    wb = load_workbook(WB, data_only=True)
    out = {"reviewer": [[c.value for c in r] for r in wb["Reviewer"].iter_rows()][0][1], "rules": {}}
    for sheet in ("Cause mentions", "Payment provisions"):
        ws = wb[sheet]; hdr = [c.value for c in ws[1]]
        ans_col = 5
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            k = key[(sheet, int(r[0]))]
            a = str(r[ans_col]).strip() if r[ans_col] is not None else ""
            if a not in ("Yes", "No", "Unclear"):
                raise SystemExit(f"{sheet} row {r[0]}: missing or invalid answer {a!r}")
            dist = nearest(r[hdr.index("Passage")], cause_pats[k["cause"]], other[sheet])
            rows.append({"flag": k["rule"] == "flagged", "ans": a, "dist": dist, "state": k["state"]})
        fl = [x for x in rows if x["flag"]]; nf = [x for x in rows if not x["flag"]]
        def ppv(sub, strict=False):
            n = [x for x in sub if strict or x["ans"] != "Unclear"]
            k = sum(x["ans"] == "Yes" for x in n)
            return {"k": k, "n": len(n), "value": round(k / len(n), 3) if n else None, "ci": wilson(k, len(n))}
        def npv(sub, strict=False):
            n = [x for x in sub if strict or x["ans"] != "Unclear"]
            k = sum(x["ans"] == "No" for x in n)
            return {"k": k, "n": len(n), "value": round(k / len(n), 3) if n else None, "ci": wilson(k, len(n))}
        narrow = [x for x in fl if x["dist"] is not None and x["dist"] <= NARROW[sheet]]
        out["rules"][sheet] = {
            "n_flagged": len(fl), "n_not_flagged": len(nf),
            "n_unclear": sum(x["ans"] == "Unclear" for x in rows),
            "n_states": len({x["state"] for x in rows}),
            "ppv": ppv(fl), "npv": npv(nf),
            "ppv_unclear_as_wrong": ppv(fl, True), "npv_unclear_as_wrong": npv(nf, True),
            "narrow_window_chars": NARROW[sheet], "ppv_narrow_window": ppv(narrow),
            "flagged_distance_unmeasurable": sum(x["dist"] is None for x in fl)}
    json.dump(out, open(RES / "contract_validation_scores.json", "w"), indent=1)
    print(json.dumps(out, indent=1))


if __name__ == "__main__":
    main()
