"""
36_dsl_inference.py — Design-based estimates of the share of recommendations that
name an intervention with evidence of benefit, correcting the automated labels
with the physician-coded labels.

Estimands. theta1, the share of all recommendations that name an intervention with
evidence of benefit; theta2, that share among recommendations naming any
intervention in the table. Truth for each recommendation is one of three
categories: an intervention with evidence of benefit, another table intervention,
or no table intervention.

Design. Physician labels come from four sources, each with a known inclusion
probability: the blinded sample (script 24; stratified by intervention and cause
within the recommendations not shown in the adjudication workbook), the second
coder's additional sample (script 34; simple random sample of recommendations the
final labels mark as naming no intervention), physician adjudication of
disagreements on the intervention field (probability 1), and the supplementary
sample of recommendations that neither sample could reach (script 37). Inclusion
probabilities for the blinded sample are computed by replicating its sampling
algorithm N_SIM times; the replication must reproduce the drawn sample at the
recorded seed.

Estimator. Design-based supervised learning (Egami et al., NeurIPS 2023) with the
automated label entering as a stratifier: recommendations are grouped into cells by
final automated category and by design group, the true-category distribution in
each cell is estimated from its physician-labeled members weighted by inverse
inclusion probability, and cell estimates are summed with cell sizes as weights.
This is the prediction-powered difference estimator (Angelopoulos et al., Science
2023) written for a stratified sample. Uncertainty is propagated by drawing each
cell's category distribution from a Dirichlet distribution with Jeffreys prior and
Kish effective sample size, because observed error counts are near zero and the
unit-level Wald interval collapses when no errors are observed; that interval is
reported alongside for comparison. Cells with no physician-labeled member are
bounded (every member assigned the category that minimizes, then maximizes, the
estimand) and, as a secondary analysis, assigned the error distribution of the
matching category in the blinded frame.

Only workbooks completed by hand count as physician labels. A full run reads a gold
file only if its name is listed in validation/GOLD_CONFIRMED.txt (one file name per
line, added by the investigator) and stops if a present file is not listed.
--dry-run tests the code on whatever files are present and writes
results/dsl_inference_DRYRUN.json, which no downstream script reads.

Output: results/dsl_inference.json
"""
from __future__ import annotations
import argparse, csv, importlib.util, json, random, sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
RES, VAL = ROOT / "results", ROOT / "validation"
SEED_BLIND, SEED_SECOND, N_BLIND, N_SECOND = 20260927, 20260930, 45, 50
N_SIM, N_DRAW, SEED = 20000, 20000, 20261004
GOLD = {"blinded": VAL / "MMRC_blinded_coding_ADJ.xlsx",
        "second": VAL / "MMRC_second_coder_CA.xlsx",
        "adjudication": VAL / "MMRC_adjudication_ADJ.xlsx",
        "supplement": VAL / "MMRC_supplementary_coding_completed.xlsx"}
CATS = ("benefit", "other", "none")
INT_COL = "Specific intervention requested"


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def jl(p, field=None):
    rows = {json.loads(l)["id"]: json.loads(l) for l in open(p)}
    return {k: (v[field] if field else v) for k, v in rows.items()}


def strata(frame_recs, I, A):
    by_int, by_cause = {}, {}
    for r in frame_recs:
        by_int.setdefault(I.get(r["id"], "none"), []).append(r["id"])
        by_cause.setdefault(A[r["id"]]["cause_domain"], []).append(r["id"])
    return [r["id"] for r in frame_recs], sorted(by_int.items()), sorted(by_cause.items())


def blinded_draw(st, A, seed):
    """Script 24's sampling algorithm; same random-number calls in the same order."""
    order, by_int, by_cause = st
    rng = random.Random(seed); chosen = []
    for k, ids in by_int:
        chosen += rng.sample(ids, min(len(ids), 2))
    cs = set(chosen)
    for c, ids in by_cause:
        have = sum(A[i]["cause_domain"] == c for i in chosen)
        pool = [i for i in ids if i not in cs]
        add = rng.sample(pool, max(0, min(len(pool), 2 - have)))
        chosen += add; cs.update(add)
    rest = [i for i in order if i not in cs]
    chosen += rng.sample(rest, max(0, N_BLIND - len(chosen)))
    rng.shuffle(chosen)
    return chosen


def read_codes(path, sheet, key_rows, lab2key):
    """Recommendation id -> intervention key from a coded workbook."""
    ws = load_workbook(path, data_only=True)[sheet]
    hdr = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        v = r[hdr.index(INT_COL)]
        v = v.strip() if isinstance(v, str) else v
        if v in (None, ""):
            raise SystemExit(f"{path.name} / {sheet}: row {r[0]} has no intervention code")
        if v not in lab2key:
            raise SystemExit(f"{path.name} / {sheet}: row {r[0]} has an unrecognized intervention: {v!r}")
        out[key_rows[int(r[0])]] = lab2key[v]
    return out


def main() -> None:
    ap = argparse.ArgumentParser(); ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--exclude", nargs="*", default=[], choices=list(GOLD),
                    help="treat these gold sources as not yet collected (interim run)")
    args = ap.parse_args()
    lc = _load("lc", "07_llm_classify.py"); iv = _load("iv", "12_interventions.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    benefit = {i.key for i in iv.build(P) if i.grade.startswith(("demonstrated", "moderate"))}
    lab2key = {i.label: i.key for i in iv._registry(P)}; lab2key["No specific intervention"] = "none"
    cat = lambda k: "benefit" if k in benefit else ("none" if k == "none" else "other")

    A = jl(RES / "llm_labels_passA.jsonl"); FI = jl(RES / "intervention_map_passA.jsonl", "intervention")
    ids = sorted(FI); N = len(ids)
    auto = {i: cat(FI[i]) for i in ids}
    present = {k: p.exists() and k not in args.exclude for k, p in GOLD.items()}
    if not args.dry_run:
        conf_file = VAL / "GOLD_CONFIRMED.txt"
        confirmed = set(conf_file.read_text().split()) if conf_file.exists() else set()
        unconfirmed = [p.name for k, p in GOLD.items() if present[k] and p.name not in confirmed]
        if unconfirmed:
            raise SystemExit(f"not confirmed as coded by hand: {unconfirmed}; list them in {conf_file.name} "
                             "or remove them")
        if not present["blinded"]:
            raise SystemExit("blinded coding workbook missing")

    # ---- inclusion probabilities ----
    adj_key = list(csv.DictReader(open(RES / "adjudication_key.csv")))
    seen = {int(r["id_or_string"]) for r in adj_key if not r["sheet"].endswith("Cause labels")}
    frame = [r for r in lc.load_recs() if r["id"] not in seen]
    vkey = [r for r in csv.DictReader(open(RES / "validation_key.csv")) if r["sheet"] == "Recommendations"]
    first = [int(r["id_or_string"]) for r in sorted(vkey, key=lambda r: int(r["row"]))]
    st = strata(frame, FI, A)
    if blinded_draw(st, A, SEED_BLIND) != first:
        raise SystemExit("replicated sampling algorithm does not reproduce the blinded sample")
    none_ids = [i for i in ids if FI[i] == "none"]
    pi_a = dict.fromkeys(ids, 0.0); pi_ab = dict.fromkeys(ids, 0.0)
    for s in range(N_SIM):
        ch = set(blinded_draw(st, A, 10_000_000 + s))
        q = N_SECOND / (len(none_ids) - sum(i in ch for i in none_ids))
        for i in ch:
            pi_a[i] += 1; pi_ab[i] += 1
        if present["second"]:
            for i in none_ids:
                if i not in ch:
                    pi_ab[i] += q
    pi = {i: pi_ab[i] / N_SIM for i in ids}
    exp_blind = sum(pi_a.values()) / N_SIM  # the stratified draw can exceed N_BLIND for some seeds
    assert not present["second"] or abs(sum(pi.values()) - exp_blind - N_SECOND) < 1e-6

    gold, source = {}, {}
    # adjudicated disagreements on the intervention field: census
    adj_int = {}
    if present["adjudication"]:
        ws_map = load_workbook(GOLD["adjudication"], data_only=True)
        for r in adj_key:
            if r["field"] != INT_COL:
                continue
            ws = ws_map[r["sheet"]] if r["sheet"] in ws_map.sheetnames else None
            if ws is None:
                continue
            hdr = [c.value for c in ws[1]]
            if ws.cell(row=int(r["excel_row"]), column=hdr.index("Your decision") + 1).value:
                adj_int[int(r["id_or_string"])] = FI[int(r["id_or_string"])]  # final label = physician's decision
    for i, k in adj_int.items():
        pi[i] = 1.0; gold[i] = cat(k); source[i] = "adjudication"
    # supplementary census and sample (script 37)
    skey_path = RES / "supplementary_coding_key.csv"
    if skey_path.exists() and present["supplement"]:
        skey = list(csv.DictReader(open(skey_path)))
        for r in skey:
            i = int(r["id"]); pi[i] = 1 - (1 - pi[i]) * (1 - float(r["inclusion_probability"]))
        rows = {int(r["row"]): int(r["id"]) for r in skey}
        for i, v in read_codes(GOLD["supplement"], "Recommendations", rows, lab2key).items():
            gold.setdefault(i, cat(v)); source.setdefault(i, "supplement")
    if present["blinded"]:
        rows = {n: i for n, i in enumerate(first, 1)}
        for i, v in read_codes(GOLD["blinded"], "Recommendations", rows, lab2key).items():
            gold[i] = cat(v); source[i] = "blinded"
    if present["second"]:
        sk = [r for r in csv.DictReader(open(RES / "second_coder_key.csv")) if r["sheet"] == "Additional recommendations"]
        rows = {int(r["row"]): int(r["id_or_string"]) for r in sk}
        for i, v in read_codes(GOLD["second"], "Additional recommendations", rows, lab2key).items():
            gold.setdefault(i, cat(v)); source.setdefault(i, "second coder")

    # ---- cells ----
    def group(i):
        if pi[i] >= 1 - 1e-12:
            return "census"
        if pi[i] == 0:
            return "unreached"
        return "blinded frame" if i not in seen else "adjudication frame"
    cells = {}
    for i in ids:
        cells.setdefault((group(i), auto[i]), []).append(i)
    est = {}
    for h, mem in cells.items():
        g = [i for i in mem if i in gold]
        w = np.array([1 / pi[i] for i in g])
        if len(g) == 0:
            est[h] = {"n": len(mem), "n_gold": 0}
            continue
        p = np.array([sum(w[j] for j, i in enumerate(g) if gold[i] == c) for c in CATS]) / w.sum()
        neff = w.sum() ** 2 / (w ** 2).sum()
        est[h] = {"n": len(mem), "n_gold": len(g), "p": p.tolist(), "n_eff": float(neff),
                  "census": h[0] == "census"}

    rng = np.random.default_rng(SEED)
    def draws(fill):
        """fill: category distribution for cells with no physician label."""
        num = np.zeros(N_DRAW); den = np.zeros(N_DRAW); pn = pd = 0.0
        for h, e in est.items():
            if e["n_gold"] == 0:
                p, D = np.asarray(fill(h)), None
            else:
                p = np.asarray(e["p"])
                D = None if e["census"] or e["n_gold"] == e["n"] else \
                    rng.dirichlet(e["n_eff"] * p + 0.5, N_DRAW)
            pn += e["n"] * p[0]; pd += e["n"] * (p[0] + p[1])
            num += e["n"] * (D[:, 0] if D is not None else p[0])
            den += e["n"] * ((D[:, 0] + D[:, 1]) if D is not None else p[0] + p[1])
        return pn / N, pn / pd, num / N, num / den
    def summ(pt1, pt2, d1, d2, lo=None, hi=None):
        q = lambda a, x: float(np.percentile(a, x))
        return {"theta1": {"est": pt1, "ci": [q(lo[2] if lo else d1, 2.5), q(hi[2] if hi else d1, 97.5)]},
                "theta2": {"est": pt2, "ci": [q(lo[3] if lo else d2, 2.5), q(hi[3] if hi else d2, 97.5)]}}
    one = lambda c: [1.0 if x == c else 0.0 for x in CATS]
    lo = draws(lambda h: one("none")); hi = draws(lambda h: one("benefit"))
    lo2 = draws(lambda h: one("other"))  # minimizes theta2
    transport = lambda h: est.get(("blinded frame", h[1]), {}).get("p", one(h[1]))
    tr = draws(transport)
    unreached = sum(e["n"] for h, e in est.items() if e["n_gold"] == 0)
    bounded = summ(lo[0], lo[1], None, None, lo, hi)
    bounded["theta2"]["ci"][0] = float(np.percentile(lo2[3], 2.5))
    bounded["theta1"]["est_range"] = [lo[0], hi[0]]; bounded["theta2"]["est_range"] = [lo2[1], hi[1]]
    bounded["theta1"].pop("est"); bounded["theta2"].pop("est")

    # unit-level prediction-powered estimate with Wald interval (reached units only)
    yhat = {i: auto[i] == "benefit" for i in ids}
    r_ids = [i for i in gold if pi[i] > 0]
    corr = np.array([(float(gold[i] == "benefit") - yhat[i]) / pi[i] for i in r_ids])
    t_ppi = (sum(yhat.values()) + corr.sum()) / N
    v_ppi = sum((1 - pi[i]) / pi[i] ** 2 * (float(gold[i] == "benefit") - yhat[i]) ** 2 for i in r_ids) / N ** 2

    out = {"gold_status": "DRY RUN — gold files not confirmed as manually coded" if args.dry_run
           else "manually coded gold labels", "gold_files_present": present,
           "n": N, "automated": {"theta1": sum(yhat.values()) / N,
                                 "theta2": sum(yhat.values()) / sum(auto[i] != "none" for i in ids)},
           "n_gold": len(gold), "gold_by_source": {s: sum(v == s for v in source.values()) for s in set(source.values())},
           "n_unreached": unreached,
           "cells": {f"{g} | {c}": e for (g, c), e in sorted(est.items())},
           "primary_bounded": bounded,
           "secondary_transport": summ(*tr),
           "ppi_unit_wald": {"theta1": t_ppi, "ci": [t_ppi - 1.96 * v_ppi ** 0.5, t_ppi + 1.96 * v_ppi ** 0.5],
                             "note": "excludes unreached units; collapses when no errors are observed"},
           "expected_blinded_sample_size": exp_blind,
           "settings": {"n_sim": N_SIM, "n_draw": N_DRAW, "seed": SEED}}
    if unreached == 0:
        out["primary"] = summ(*tr)
    out["excluded_sources"] = args.exclude
    fn = RES / ("dsl_inference_DRYRUN.json" if args.dry_run else
                "dsl_inference_INTERIM.json" if args.exclude else "dsl_inference.json")
    json.dump(out, open(fn, "w"), indent=1)
    print(json.dumps({k: out[k] for k in ("gold_status", "automated", "gold_by_source", "n_unreached",
                                          "primary_bounded", "secondary_transport", "ppi_unit_wald")}, indent=1))
    for h, e in out["cells"].items():
        print(f"  {h:45s} n={e['n']:5d} gold={e['n_gold']:3d} p={np.round(e.get('p', []), 3).tolist()}")


if __name__ == "__main__":
    main()
