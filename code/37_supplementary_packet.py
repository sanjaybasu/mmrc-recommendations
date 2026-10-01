"""
37_supplementary_packet.py — Physician coding of the recommendations that neither
the blinded sample nor the second coder's sample could reach.

The blinded sample (script 24) was drawn from recommendations not shown in the
adjudication workbook, and the second coder's additional sample (script 34) from
recommendations the final labels mark as naming no intervention. A recommendation
shown in the adjudication workbook whose final label names an intervention, and
whose intervention label was not adjudicated, therefore had no chance of physician
review. Design-based estimation (script 36) needs every recommendation to have a
known, nonzero chance. This packet covers them: every such recommendation whose
final label or either model's label is an intervention with evidence of benefit
(census), and a simple random sample of N_SAMPLE of the rest.

Only the intervention field is coded. No automated label appears in the workbook.

Output: validation/MMRC_supplementary_coding.xlsx, results/supplementary_coding_key.csv
"""
from __future__ import annotations
import csv, importlib.util, json, random, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RES, OUT = ROOT / "results", ROOT / "validation"
SEED, N_SAMPLE = 20261003, 25
FILL = PatternFill("solid", fgColor="DDE6EE")
INT_COL = "Specific intervention requested"


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def main() -> None:
    lc = _load("lc", "07_llm_classify.py"); iv = _load("iv", "12_interventions.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    benefit = {i.key for i in iv.build(P) if i.grade.startswith(("demonstrated", "moderate"))}
    int_options = [i.label for i in iv._registry(P)] + ["No specific intervention"]
    im = lambda f: {json.loads(l)["id"]: json.loads(l)["intervention"] for l in open(RES / f)}
    FI, I1, I2 = im("intervention_map_passA.jsonl"), im("imap_claude-opus-5-5.jsonl"), im("imap_gpt-6-astra.jsonl")
    key = list(csv.DictReader(open(RES / "adjudication_key.csv")))
    seen = {int(r["id_or_string"]) for r in key if not r["sheet"].endswith("Cause labels")}
    adjudicated = {int(r["id_or_string"]) for r in key if r["sheet"] == "1 Priority" and r["field"] == INT_COL}
    unreached = sorted(i for i in seen if FI[i] != "none" and i not in adjudicated)
    census = [i for i in unreached if {FI[i], I1[i], I2[i]} & benefit]
    rest = [i for i in unreached if i not in set(census)]
    sample = random.Random(SEED).sample(rest, min(N_SAMPLE, len(rest)))
    items = [(i, 1.0) for i in census] + [(i, len(sample) / len(rest)) for i in sample]
    random.Random(SEED + 1).shuffle(items)

    recs = {r["id"]: r for r in lc.load_recs()}
    wb = Workbook(); ws = wb.active; ws.title = "Recommendations"
    cols = [("Row", 6), ("State", 7), ("Report section or category", 28), ("Recommendation text", 80),
            (INT_COL, 42), ("Unsure? (Y)", 10), ("Comment", 30)]
    for j, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[c.column_letter].width = w
    lst = wb.create_sheet("list"); lst.sheet_state = "hidden"
    for j, v in enumerate(int_options, 1):
        lst.cell(row=j, column=1, value=v)
    dv = DataValidation(type="list", formula1=f"'list'!$A$1:$A${len(int_options)}", allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"E2:E{len(items) + 1}")
    for n, (i, _) in enumerate(items, 2):
        r = recs[i]
        for j, v in enumerate([n - 1, r["state"], r["category"], r["text"]], 1):
            ws.cell(row=n, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    ws2 = wb.create_sheet("Coder"); ws2["A1"] = "Coder initials"; ws2["A2"] = "Date completed"
    wb.save(OUT / "MMRC_supplementary_coding.xlsx")
    with open(RES / "supplementary_coding_key.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["row", "id", "inclusion_probability"])
        w.writerows([[n, i, p] for n, (i, p) in enumerate(items, 1)])
    print(f"unreached {len(unreached)}: census {len(census)}, sample {len(sample)} of {len(rest)}")


if __name__ == "__main__":
    main()
