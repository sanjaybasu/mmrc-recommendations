"""
34_second_coder_packet.py — Blinded coding workbook for a second physician.

Sheet "Recommendations" repeats the 45 recommendations and sheet "Cause labels" the
15 cause labels the first physician coded, in the same order, so that agreement
between the 2 physicians can be computed. Sheet "Additional recommendations" holds
50 recommendations drawn at random from those the final labels mark as naming no
intervention in the table, to estimate how often an intervention was missed.
No automated labels appear anywhere in the workbook.

Output: validation/MMRC_second_coder.xlsx, results/second_coder_key.csv
"""
from __future__ import annotations
import csv, importlib.util, json, random, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RES, OUT = ROOT / "results", ROOT / "validation"
SEED, N_EXTRA = 20260930, 50
FILL = PatternFill("solid", fgColor="DDE6EE")


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def main() -> None:
    lc = _load("lc", "07_llm_classify.py"); iv = _load("iv", "12_interventions.py")
    P = json.load(open(RES / "model_params.json"))["params"]
    int_options = [i.label for i in iv._registry(P)] + ["No specific intervention"]
    recs = {r["id"]: r for r in lc.load_recs()}
    key = list(csv.DictReader(open(RES / "validation_key.csv")))
    first = [int(r["id_or_string"]) for r in sorted((r for r in key if r["sheet"] == "Recommendations"),
                                                    key=lambda r: int(r["row"]))]
    causes = [r["id_or_string"] for r in sorted((r for r in key if r["sheet"] == "Cause labels"),
                                                key=lambda r: int(r["row"]))]
    final = {json.loads(l)["id"]: json.loads(l)["intervention"] for l in open(RES / "intervention_map_passA.jsonl")}
    pool = [i for i, v in final.items() if v == "none" and i not in set(first)]
    extra = random.Random(SEED).sample(pool, N_EXTRA)

    cols = [("Row", 6), ("State", 7), ("Report section or category", 28), ("Recommendation text", 70),
            ("Cause of death addressed", 26), ("Policy lever", 26), ("Most authoritative actor named", 30),
            ("Coverage period", 22), ("Specific intervention requested", 38), ("Unsure? (Y)", 10),
            ("Comment", 30)]
    options = {5: lc.CAUSE_DOMAINS, 6: lc.LEVERS, 7: lc.AUTHORITY, 8: lc.PERIODS, 9: int_options}
    wb = Workbook()
    lists = {}

    def hidden_list(name, values):
        ws = wb.create_sheet(name)
        for j, v in enumerate(values, 1):
            ws.cell(row=j, column=1, value=v)
        ws.sheet_state = "hidden"
        return f"'{name}'!$A$1:$A${len(values)}"

    for col, vals in options.items():
        lists[col] = hidden_list(f"list{col}", vals)
    lists["cause"] = hidden_list("listcause", [c for c in lc.CAUSE_DOMAINS if c != "Cross-cutting"]
                                 + ["Not a cause of death"])

    def rec_sheet(ws, ids):
        for j, (h, w) in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[c.column_letter].width = w
        for n, i in enumerate(ids, 2):
            r = recs[i]
            for j, v in enumerate([n - 1, r["state"], r["category"], r["text"]], 1):
                ws.cell(row=n, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
        for col in options:
            dv = DataValidation(type="list", formula1=lists[col], allow_blank=True)
            ws.add_data_validation(dv)
            letter = ws.cell(row=1, column=col).column_letter
            dv.add(f"{letter}2:{letter}{len(ids) + 1}")
        ws.freeze_panes = "E2"

    ws1 = wb.active; ws1.title = "Recommendations"
    rec_sheet(ws1, first)
    ws2 = wb.create_sheet("Additional recommendations", 1)
    rec_sheet(ws2, extra)
    ws3 = wb.create_sheet("Cause labels", 2)
    for j, (h, w) in enumerate([("Row", 6), ("Committee-written cause label", 60),
                                ("Cause category", 30), ("Comment", 30)], 1):
        c = ws3.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL
        ws3.column_dimensions[c.column_letter].width = w
    for n, s in enumerate(causes, 2):
        ws3.cell(row=n, column=1, value=n - 1); ws3.cell(row=n, column=2, value=s)
    dv = DataValidation(type="list", formula1=lists["cause"], allow_blank=True)
    ws3.add_data_validation(dv); dv.add(f"C2:C{len(causes) + 1}")
    ws4 = wb.create_sheet("Coder", 3)
    ws4["A1"] = "Coder initials"; ws4["A2"] = "Date completed"
    wb.save(OUT / "MMRC_second_coder.xlsx")

    with open(RES / "second_coder_key.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["sheet", "row", "id_or_string"])
        w.writerows([["Recommendations", n, i] for n, i in enumerate(first, 1)])
        w.writerows([["Additional recommendations", n, i] for n, i in enumerate(extra, 1)])
        w.writerows([["Cause labels", n, s] for n, s in enumerate(causes, 1)])
    print(f"second coder packet: {len(first)} + {len(extra)} recommendations, {len(causes)} cause labels")


if __name__ == "__main__":
    main()
