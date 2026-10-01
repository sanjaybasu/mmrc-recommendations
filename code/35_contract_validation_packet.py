"""
35_contract_validation_packet.py — Blinded review of the text rules that decide
whether a Medicaid contract names a cause of pregnancy-related death, and whether it
ties that cause to a payment consequence.

Sheet "Cause mentions": 80 passages, each containing a cause term. Half were
flagged by the rule (a maternal or perinatal term within 300 characters) and half
were not. The reviewer judges whether the passage refers to that condition in
pregnant or postpartum people. This estimates the rule's positive and negative
predictive values.

Sheet "Payment provisions": 80 passages, each a maternal-context mention of a cause.
Half were flagged as tied to a withhold or incentive (payment language within
1,000 characters) and half were not. The reviewer judges whether the passage ties
care for that condition in pregnant or postpartum people to a payment consequence.

Passages are sampled across states and causes with a fixed seed and shuffled;
the flag is recorded only in results/contract_validation_key.csv.

Output: validation/MMRC_contract_text_validation.xlsx, results/contract_validation_key.csv
"""
from __future__ import annotations
import bisect, csv, importlib.util, random, re, sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
RES, OUT = ROOT / "results", ROOT / "validation"
SEED, N_EACH, DOCS_PER_STATE = 20261001, 40, 12
FILL = PatternFill("solid", fgColor="DDE6EE")


def _load(n, f):
    s = importlib.util.spec_from_file_location(n, ROOT / "code" / f)
    m = importlib.util.module_from_spec(s); sys.modules[n] = m; s.loader.exec_module(m)
    return m


def near(pos, spans, w):
    j = bisect.bisect_left(spans, pos)
    return (j < len(spans) and spans[j] - pos <= w) or (j > 0 and pos - spans[j - 1] <= w)


def snippet(text, a, b, pad):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    s, e = max(0, a - pad), min(len(text), b + pad)
    body = ILLEGAL_CHARACTERS_RE.sub(" ", re.sub(r"\s+", " ", text[s:e])).strip()
    return ("…" if s > 0 else "") + body + ("…" if e < len(text) else "")


def main() -> None:
    chain = _load("chain", "04_accountability_chain.py")
    rng = random.Random(SEED)
    pats = {d: re.compile(p, re.I) for d, p in chain.CAUSE_TERMS.items()}
    buckets = {"cause_flag": [], "cause_noflag": [], "pay_flag": [], "pay_noflag": []}
    for name, abbr in chain.STATE_ABBR.items():
        sd = chain.RFP / name
        if not sd.exists():
            continue
        files = sorted(sd.rglob("*.txt")); rng.shuffle(files)
        for f in files[:DOCS_PER_STATE]:
            text = f.read_text(errors="ignore")
            mat = [m.start() for m in chain.MATERNAL_CONTEXT.finditer(text)]
            pay = [m.start() for m in chain.WITHHOLD.finditer(text)]
            for d, pat in pats.items():
                hits = list(pat.finditer(text))
                rng.shuffle(hits)
                for m in hits[:6]:
                    i = m.start()
                    rec = {"state": abbr, "doc": f.name, "cause": d}
                    if mat and near(i, mat, chain.CONTEXT_WINDOW):
                        buckets["cause_flag"].append({**rec, "text": snippet(text, m.start(), m.end(), 350)})
                        if pay and near(i, pay, 1000):
                            buckets["pay_flag"].append({**rec, "text": snippet(text, m.start(), m.end(), 600)})
                        else:
                            buckets["pay_noflag"].append({**rec, "text": snippet(text, m.start(), m.end(), 600)})
                    else:
                        buckets["cause_noflag"].append({**rec, "text": snippet(text, m.start(), m.end(), 350)})
        print(f"  {name}", flush=True)

    def draw(items, n):
        """Spread the draw across states and causes."""
        by = {}
        for it in items:
            by.setdefault((it["state"], it["cause"]), []).append(it)
        keys = list(by); rng.shuffle(keys)
        out = []
        while len(out) < n and any(by.values()):
            for k in keys:
                if by[k] and len(out) < n:
                    out.append(by[k].pop(rng.randrange(len(by[k]))))
        return out

    sheets = {"Cause mentions": [("flagged", x) for x in draw(buckets["cause_flag"], N_EACH)]
              + [("not flagged", x) for x in draw(buckets["cause_noflag"], N_EACH)],
              "Payment provisions": [("flagged", x) for x in draw(buckets["pay_flag"], N_EACH)]
              + [("not flagged", x) for x in draw(buckets["pay_noflag"], N_EACH)]}
    question = {"Cause mentions": "Does this passage refer to this condition in pregnant or postpartum people?",
                "Payment provisions": "Does this passage tie care for this condition in pregnant or postpartum people to a payment consequence?"}
    wb = Workbook(); wb.remove(wb.active)
    key = []
    for sname, items in sheets.items():
        rng.shuffle(items)
        ws = wb.create_sheet(sname)
        cols = [("Row", 6), ("State", 7), ("Document", 30), ("Condition", 24), ("Passage", 100),
                (question[sname], 22), ("Comment", 30)]
        for j, (h, w) in enumerate(cols, 1):
            c = ws.cell(row=1, column=j, value=h); c.font = Font(bold=True); c.fill = FILL
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[c.column_letter].width = w
        dv = DataValidation(type="list", formula1='"Yes,No,Unclear"', allow_blank=True)
        ws.add_data_validation(dv)
        for n, (flag, it) in enumerate(items, 2):
            for j, v in enumerate([n - 1, it["state"], it["doc"], it["cause"], it["text"]], 1):
                ws.cell(row=n, column=j, value=v).alignment = Alignment(wrap_text=True, vertical="top")
            dv.add(f"F{n}")
            key.append([sname, n - 1, flag, it["state"], it["cause"], it["doc"]])
        ws.freeze_panes = "F2"
    ws = wb.create_sheet("Reviewer"); ws["A1"] = "Initials"; ws["A2"] = "Date completed"
    wb.save(OUT / "MMRC_contract_text_validation.xlsx")
    with open(RES / "contract_validation_key.csv", "w", newline="") as f:
        w = csv.writer(f); w.writerow(["sheet", "row", "rule", "state", "cause", "document"])
        w.writerows(key)
    print({k: len(v) for k, v in buckets.items()}, {k: len(v) for k, v in sheets.items()})


if __name__ == "__main__":
    main()
